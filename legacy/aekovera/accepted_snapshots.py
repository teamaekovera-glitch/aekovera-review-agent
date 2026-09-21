"""Accepted companies workbook - the corrected record, exactly as uploaded (v12.28).

What goes in
------------
ONLY suppliers that were actually accepted (Platform ready confirmed). For each
one, the full company record as it stood on the review page immediately
BEFORE Platform ready was clicked - i.e. after every correction had been saved.

v32.1 ORDER (this is the fix for "accepted_companies shows old data"):
    corrections saved -> page RELOADED -> every field read from the fresh page
    -> row written here with verdict_status "pending" -> verdict pressed
    -> confirm_pending() marks it "accepted", or rollback_pending() removes it
       (or restores the company's previous accepted row).
Edits go over HTTP and never refresh the browser tab, so before v32.1 the
"re-read" saw the page as it was BEFORE the corrections.

Not written: REJECT, MANUAL_REVIEW, Auto-Mode holds, or a Platform ready that
did not land. main.py only calls record_accepted() after the accept is
confirmed, so nothing here needs to second-guess the decision.

How this differs from supplier_history.xlsx
-------------------------------------------
The history workbook's Accepted sheet RECONSTRUCTS the final values (original
record + the corrections we tried to apply) and carries the whole decision
audit block. This workbook holds the record READ BACK from the page, so it
reflects what the platform actually has, and it is laid out as a clean company
list: company fields first, a short edit summary at the end.

Layout
------
One sheet, one row per company keyed on the MASTER id (history.dedupe_key).
Accepting the same company again refreshes its row in place: first_accepted
is kept, times_accepted goes up. Company-field columns are discovered per
record, so a new field gets a new column instead of being dropped.

The sheet is rebuilt on every save so columns stay in a stable order
(identity -> company fields -> edit summary) no matter when a field first
appeared. Any other sheets in the file are left untouched.

Durability
----------
Saved after every accepted company. If the file is open in Excel (read-only
on Windows) or openpyxl is missing, the row is queued to
ACCEPTED_SNAPSHOT_PENDING_FILE and merged on the next successful save.
Never raises - a record-keeping problem must not stop a review session.
"""

import json
import os
from datetime import datetime
from pathlib import Path

from config import (
    ENABLE_ACCEPTED_SNAPSHOT,
    ACCEPTED_SNAPSHOT_FILE,
    ACCEPTED_SNAPSHOT_SHEET,
    ACCEPTED_SNAPSHOT_PENDING_FILE,
)
from history import DETAIL_SKIP, PREFERRED_DETAIL_ORDER, _text, dedupe_key

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None
    load_workbook = None


LEAD_COLUMNS = ["accepted_at", "record_id", "company_name"]
TAIL_COLUMNS = [
    "verdict_status",
    "edit_status",
    "edits_made",
    "fields_failed",
    "fields_unresolved",
    "fields_not_landed",
    "confirmed_by",
    "snapshot_source",
    "first_accepted",
    "times_accepted",
]
_FIXED = set(LEAD_COLUMNS) | set(TAIL_COLUMNS)

SOURCE_PAGE = "page reloaded and re-read after corrections, before the verdict"
SOURCE_RECONSTRUCTED = "reconstructed (page re-read failed)"

_WIDE = {"edits_made", "description", "products", "specialty", "address"}
_HEADER_FILL = "FFEFEFEF"
_DIFF_VALUE_LIMIT = 120


# ---------------------------------------------------------------------------
# Row construction
# ---------------------------------------------------------------------------

def _short(value):
    text = _text(value)
    if len(text) > _DIFF_VALUE_LIMIT:
        text = text[: _DIFF_VALUE_LIMIT - 1] + "…"
    return text


def _reconstruct(record, result, applied, cleared):
    """Fallback when the page could not be re-read: original + saved changes."""
    fields = dict(record.get("fields") or {})
    applied_set = {_text(f) for f in applied or ()}
    for change in result.get("changes") or []:
        if not isinstance(change, dict):
            continue
        field = _text(change.get("field"))
        if field in applied_set and change.get("new_value") is not None:
            fields[field] = change.get("new_value")
    for item in cleared or ():
        fields[_text(item[0])] = ""
    return fields


def _diff(before, after):
    """'field: old -> new' for every field whose value changed during the run."""
    lines = []
    for key in list(before) + [k for k in after if k not in before]:
        if _text(key) in DETAIL_SKIP and _text(key) != "company_name":
            continue
        old, new = _text(before.get(key)), _text(after.get(key))
        if old == new:
            continue
        if not new:
            lines.append(f"{key}: cleared (was {_short(old)!r})")
        elif not old:
            lines.append(f"{key}: added {_short(new)!r}")
        else:
            lines.append(f"{key}: {_short(old)!r} -> {_short(new)!r}")
    return lines


def build_row(record, corrected, result, *, applied=(), cleared=(), failed=(),
              needs_clear=(), needs_review=(), confirmed_by=""):
    original = dict(record.get("fields") or {})
    if corrected and corrected.get("fields"):
        fields = dict(corrected["fields"])
        record_id = _text(corrected.get("record_id")) or _text(record.get("record_id"))
        source = SOURCE_PAGE
    else:
        fields = _reconstruct(record, result, applied, cleared)
        record_id = _text(record.get("record_id"))
        source = SOURCE_RECONSTRUCTED

    edits = _diff(original, fields)
    failed = [_text(f) for f in failed or () if _text(f)]
    unresolved = [_text(x[0]) for x in list(needs_clear or ()) + list(needs_review or ())]
    not_landed = list((corrected or {}).get("not_landed") or [])

    if failed or unresolved or not_landed:
        parts = []
        if not_landed:
            parts.append(f"{len(not_landed)} not on the reloaded record")
        if failed:
            parts.append(f"{len(failed)} failed")
        if unresolved:
            parts.append(f"{len(unresolved)} unresolved")
        status = "partial - " + ", ".join(parts)
    elif edits:
        status = "complete"
    else:
        status = "no changes needed"

    now = datetime.now().isoformat(timespec="seconds")
    row = {
        "accepted_at": now,
        "record_id": record_id,
        "company_name": _text(fields.get("company_name")) or _text(result.get("company_name")),
        "edit_status": status,
        "edits_made": "\n".join(edits),
        "fields_failed": ", ".join(failed),
        "fields_unresolved": ", ".join(unresolved),
        "fields_not_landed": "; ".join(
            f"{f}: saved {_short(e)!r}, page shows {_short(p)!r}" for f, e, p in not_landed),
        "verdict_status": "accepted",
        "confirmed_by": _text(confirmed_by),
        "snapshot_source": source,
        "first_accepted": now,
        "times_accepted": 1,
    }
    for key, value in fields.items():
        key = _text(key)
        if key and key not in DETAIL_SKIP and key not in _FIXED:
            row[key] = _text(value)
    return row


# ---------------------------------------------------------------------------
# Pending spillover
# ---------------------------------------------------------------------------

def _append_pending(row):
    try:
        with Path(ACCEPTED_SNAPSHOT_PENDING_FILE).open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
        return True
    except Exception as exc:  # pragma: no cover - disk-level failure
        print(f"⚠ Could not write {ACCEPTED_SNAPSHOT_PENDING_FILE} either: {exc}")
        return False


def _read_pending():
    path = Path(ACCEPTED_SNAPSHOT_PENDING_FILE)
    if not path.exists():
        return []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rows.append(json.loads(line))
        except Exception:
            continue
    return rows


def _clear_pending():
    try:
        Path(ACCEPTED_SNAPSHOT_PENDING_FILE).unlink()
    except FileNotFoundError:
        pass
    except Exception as exc:  # pragma: no cover
        print(f"⚠ Could not clear {ACCEPTED_SNAPSHOT_PENDING_FILE}: {exc}")


# ---------------------------------------------------------------------------
# Workbook I/O
# ---------------------------------------------------------------------------

def _open_workbook():
    path = Path(ACCEPTED_SNAPSHOT_FILE)
    if path.exists():
        try:
            return load_workbook(path)
        except Exception as exc:
            # Never silently overwrite a file we cannot read - move it aside.
            backup = path.with_name(
                f"{path.stem}.corrupt-{datetime.now():%Y%m%d-%H%M%S}{path.suffix}"
            )
            try:
                path.rename(backup)
                print(f"⚠ {ACCEPTED_SNAPSHOT_FILE} could not be opened ({exc}); "
                      f"moved it to {backup.name} and started a fresh one.")
            except Exception:
                raise
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_text(h) for h in rows[0]]
    data = []
    for values in rows[1:]:
        if not any(v not in (None, "") for v in values):
            continue
        data.append({h: v for h, v in zip(headers, values) if h})
    return headers, data


def _upsert(rows, index, row):
    key = dedupe_key(row.get("record_id"), row.get("company_name"))
    position = index.get(key) if key else None
    if position is None:
        rows.append(row)
        if key:
            index[key] = len(rows) - 1
        return "added"
    old = rows[position]
    row["first_accepted"] = _text(old.get("first_accepted")) or row["first_accepted"]
    try:
        row["times_accepted"] = int(old.get("times_accepted") or 0) + 1
    except (TypeError, ValueError):
        row["times_accepted"] = 2
    # Keep any column the old row had that this read did not expose (e.g. a
    # note typed into the sheet by hand) instead of blanking it.
    merged = dict(old)
    merged.update(row)
    rows[position] = merged
    return "updated"


def _column_order(existing_headers, rows):
    """Identity, then company fields, then the edit summary.

    Well-known fields (email, phone, website, address...) always sit in their
    natural order; anything else keeps the position it already had in the
    sheet, and brand-new fields are appended alphabetically after those.
    """
    existing = [h for h in existing_headers if h and h not in _FIXED]
    seen = set(existing)
    new = sorted({k for r in rows for k in r if k and k not in _FIXED and k not in seen})
    arrival = {name: i for i, name in enumerate(existing + new)}

    def sort_key(name):
        if name in PREFERRED_DETAIL_ORDER:
            return (0, PREFERRED_DETAIL_ORDER.index(name))
        return (1, arrival[name])

    details = sorted(existing + new, key=sort_key)
    return LEAD_COLUMNS + details + TAIL_COLUMNS


def _write_sheet(wb, headers, rows):
    position = None
    if ACCEPTED_SNAPSHOT_SHEET in wb.sheetnames:
        position = wb.sheetnames.index(ACCEPTED_SNAPSHOT_SHEET)
        wb.remove(wb[ACCEPTED_SNAPSHOT_SHEET])
    ws = wb.create_sheet(ACCEPTED_SNAPSHOT_SHEET, position)

    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
        # A value beginning with "=" would be stored as a formula.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "D2"  # accepted_at / record_id / company_name stay visible
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, header in enumerate(headers, start=1):
        letter = get_column_letter(col)
        if header in _WIDE:
            ws.column_dimensions[letter].width = 48
            for cell in ws[letter][1:]:
                cell.alignment = wrap
        elif header == "company_name":
            ws.column_dimensions[letter].width = 32
        else:
            ws.column_dimensions[letter].width = max(12, min(len(header) + 6, 28))
    return ws


def _save(wb):
    """Write via a temp file so a crash mid-save cannot corrupt the workbook."""
    path = Path(ACCEPTED_SNAPSHOT_FILE)
    tmp = path.with_name(f".{path.name}.tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except Exception:
        try:
            tmp.unlink()
        except Exception:
            pass
        raise


def record_accepted(record, corrected, result, *, applied=(), cleared=(), failed=(),
                    needs_clear=(), needs_review=(), confirmed_by=""):
    """File one ACCEPTED company. Call only after Platform ready is confirmed."""
    if not ENABLE_ACCEPTED_SNAPSHOT:
        return
    try:
        row = build_row(
            record, corrected, result, applied=applied, cleared=cleared,
            failed=failed, needs_clear=needs_clear, needs_review=needs_review,
            confirmed_by=confirmed_by,
        )
    except Exception as exc:  # pragma: no cover - defensive
        print(f"⚠ Could not build the accepted-company row: {exc}")
        return

    name = row["company_name"] or row["record_id"] or "company"

    if load_workbook is None:
        _append_pending(row)
        print(f"⚠ openpyxl is not installed, so {ACCEPTED_SNAPSHOT_FILE} was not updated.\n"
              f"  Queued in {ACCEPTED_SNAPSHOT_PENDING_FILE}; run  pip install openpyxl")
        return

    try:
        wb = _open_workbook()
        headers, rows = ([], [])
        if ACCEPTED_SNAPSHOT_SHEET in wb.sheetnames:
            headers, rows = _read_sheet(wb[ACCEPTED_SNAPSHOT_SHEET])
        index = {}
        for i, existing in enumerate(rows):
            key = dedupe_key(existing.get("record_id"), existing.get("company_name"))
            if key:
                index[key] = i

        # Earlier rows stranded by a locked file go in first, in order.
        pending = _read_pending()
        for queued in pending:
            _upsert(rows, index, queued)
        verb = _upsert(rows, index, row)

        _write_sheet(wb, _column_order(headers, rows), rows)
        _save(wb)
        if pending:
            _clear_pending()
            print(f"  (merged {len(pending)} queued accepted compan"
                  f"{'y' if len(pending) == 1 else 'ies'} from {ACCEPTED_SNAPSHOT_PENDING_FILE})")
        note = "" if row["snapshot_source"] == SOURCE_PAGE else " [reconstructed - page re-read failed]"
        print(f"✓ Accepted company {verb}: '{name}' → {ACCEPTED_SNAPSHOT_FILE} "
              f"({len(rows)} total){note}")
    except PermissionError:
        if _append_pending(row):
            print(f"⚠ {ACCEPTED_SNAPSHOT_FILE} is locked (probably open in Excel).\n"
                  f"  '{name}' was queued in {ACCEPTED_SNAPSHOT_PENDING_FILE} and will be\n"
                  "  merged in automatically on the next accept once you close it.")
    except Exception as exc:
        if _append_pending(row):
            print(f"⚠ Could not update {ACCEPTED_SNAPSHOT_FILE} ({exc}); "
                  f"queued in {ACCEPTED_SNAPSHOT_PENDING_FILE} instead.")


# ---------------------------------------------------------------------------
# v32.1 - write BEFORE the verdict, confirm or roll back AFTER
# ---------------------------------------------------------------------------

PENDING_STATUS = "pending - written before Platform ready"
ACCEPTED_STATUS = "accepted"

# key -> {"previous": row or None, "row": row, "saved": bool}
_PENDING = {}


def _with_sheet(mutate):
    """Open the workbook, let `mutate(rows, index)` change the rows, save.

    Returns whatever mutate returns. Raises on I/O problems so callers can
    decide what a failure means.
    """
    wb = _open_workbook()
    headers, rows = ([], [])
    if ACCEPTED_SNAPSHOT_SHEET in wb.sheetnames:
        headers, rows = _read_sheet(wb[ACCEPTED_SNAPSHOT_SHEET])
    index = {}
    for i, existing in enumerate(rows):
        k = dedupe_key(existing.get("record_id"), existing.get("company_name"))
        if k:
            index[k] = i
    outcome = mutate(rows, index)
    _write_sheet(wb, _column_order(headers, rows), rows)
    _save(wb)
    return outcome


def record_pending(record, corrected, result, *, applied=(), cleared=(), failed=(),
                   needs_clear=(), needs_review=()):
    """Write the reloaded record BEFORE the verdict. Returns a key, or None."""
    if not ENABLE_ACCEPTED_SNAPSHOT:
        return None
    try:
        row = build_row(record, corrected, result, applied=applied, cleared=cleared,
                        failed=failed, needs_clear=needs_clear, needs_review=needs_review,
                        confirmed_by="")
    except Exception as exc:  # pragma: no cover
        print(f"\u26a0 Could not build the accepted-company row: {exc}")
        return None
    row["verdict_status"] = PENDING_STATUS
    key = dedupe_key(row.get("record_id"), row.get("company_name"))
    if not key:
        return None
    entry = {"previous": None, "row": row, "saved": False}
    _PENDING[key] = entry
    if load_workbook is None:
        return key                                  # filed on confirm via the queue

    def mutate(rows, index):
        pos = index.get(key)
        entry["previous"] = dict(rows[pos]) if pos is not None else None
        return _upsert(rows, index, dict(row))
    try:
        _with_sheet(mutate)
        entry["saved"] = True
        name = row.get("company_name") or row.get("record_id")
        print(f"\u2713 Reloaded record written to {ACCEPTED_SNAPSHOT_FILE} for '{name}' "
              f"(pending until Platform ready lands)")
    except PermissionError:
        print(f"\u26a0 {ACCEPTED_SNAPSHOT_FILE} is locked (open in Excel?); the row will be "
              "filed after the verdict instead.")
    except Exception as exc:
        print(f"\u26a0 Could not write the pending row ({exc}); it will be filed after the verdict.")
    return key


def confirm_pending(key, confirmed_by=""):
    """Platform ready landed: mark the pending row accepted."""
    entry = _PENDING.pop(key, None)
    if entry is None:
        return
    row = dict(entry["row"], verdict_status=ACCEPTED_STATUS, confirmed_by=_text(confirmed_by))
    if not entry["saved"] or load_workbook is None:
        # Never made it into the file: use the normal durable path (queues if locked).
        _file_row(row)
        return

    def mutate(rows, index):
        pos = index.get(key)
        if pos is None:
            _upsert(rows, index, row)
            return
        rows[pos]["verdict_status"] = ACCEPTED_STATUS
        rows[pos]["confirmed_by"] = row["confirmed_by"]
    try:
        _with_sheet(mutate)
        print(f"\u2713 Accepted company confirmed in {ACCEPTED_SNAPSHOT_FILE}: "
              f"'{row.get('company_name') or row.get('record_id')}'")
    except Exception as exc:
        # The pending row is in the file; queue a corrected copy so the status is fixed later.
        _append_pending(row)
        print(f"\u26a0 Could not mark the row accepted ({exc}); queued in "
              f"{ACCEPTED_SNAPSHOT_PENDING_FILE}.")


def rollback_pending(key, reason=""):
    """The company was NOT accepted: remove the pending row, or restore the old one."""
    entry = _PENDING.pop(key, None)
    if entry is None or not entry["saved"] or load_workbook is None:
        return

    def mutate(rows, index):
        pos = index.get(key)
        if pos is None:
            return
        if entry["previous"] is not None:
            rows[pos] = entry["previous"]
        else:
            del rows[pos]
    try:
        _with_sheet(mutate)
        print(f"\u21b7 Pending row removed from {ACCEPTED_SNAPSHOT_FILE} "
              f"({reason or 'not accepted'}).")
    except Exception as exc:
        print(f"\u26a0 Could not remove the pending row ({exc}). Open {ACCEPTED_SNAPSHOT_FILE} "
              f"and delete the row whose verdict_status is '{PENDING_STATUS}'.")


def _file_row(row):
    """Durable write of a complete row (same fallbacks as record_accepted)."""
    name = row.get("company_name") or row.get("record_id") or "company"
    if load_workbook is None:
        _append_pending(row)
        return
    try:
        pending = _read_pending()

        def mutate(rows, index):
            for queued in pending:
                _upsert(rows, index, queued)
            return _upsert(rows, index, row)
        verb = _with_sheet(mutate)
        if pending:
            _clear_pending()
        print(f"\u2713 Accepted company {verb}: '{name}' \u2192 {ACCEPTED_SNAPSHOT_FILE}")
    except Exception as exc:
        if _append_pending(row):
            print(f"\u26a0 Could not update {ACCEPTED_SNAPSHOT_FILE} ({exc}); queued in "
                  f"{ACCEPTED_SNAPSHOT_PENDING_FILE}.")
