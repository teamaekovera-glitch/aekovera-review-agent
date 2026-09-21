"""Export-parity tests: store exports match the v34 shapes exactly.

Every expected column list below is transcribed verbatim from the frozen
archive (legacy/aekovera/history.py, accepted_snapshots.py, main.py,
config.py) - NOT imported from review_hub.store.rows, so an accidental edit
to the port's constants fails here and forces a conscious decision.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from review_hub.store.exports import (
    accepted_export,
    export_discovery_failures,
    export_field_holds,
    export_manual_reviews,
    export_non_us_origins,
    export_website_checks,
    history_export,
    write_accepted_workbook,
    write_history_workbook,
)
from review_hub.store.repository import ReviewStore

# --- legacy/aekovera/history.py lines 81-107 (verbatim) --------------------
LEGACY_HISTORY_META_COLUMNS = [
    "first_seen",
    "last_seen",
    "times_seen",
    "record_id",
    "company_name",
    "decision",
    "outcome",
    "run_mode",
    "backend",
    "confidence",
    "reason",
    "scope_match",
    "food_beverage_connection",
    "qualifying_supplier_types",
    "supply_country",
    "is_us_based",
    "supply_origin_note",
    "manual_review_reason",
    "fields_applied",
    "fields_created",
    "fields_cleared",
    "fields_failed",
    "fields_left_uncleared",
    "fields_held_for_review",
    "identity_renamed",
]

# --- legacy/aekovera/accepted_snapshots.py lines 70-82 (verbatim) ----------
LEGACY_ACCEPTED_LEAD_COLUMNS = ["accepted_at", "record_id", "company_name"]
LEGACY_ACCEPTED_TAIL_COLUMNS = [
    "verdict_status",
    "edit_status",
    "edits_made",
    "fields_failed",
    "fields_unresolved",
    "fields_not_landed",
    "confirmed_by",
    "snapshot_source",
    "first_accepted",
    "times_accepted",
]

# --- legacy/aekovera/main.py CSV writer blocks (verbatim) ------------------
LEGACY_NON_US_HEADER = [
    "timestamp",
    "company_name",
    "supply_country",
    "supply_origin_note",
    "supplier_types",
]
LEGACY_MANUAL_REVIEW_HEADER = [
    "timestamp",
    "record_id",
    "company_name",
    "manual_review_reason",
    "food_beverage_connection",
    "qualifying_supplier_types",
]
LEGACY_DISCOVERY_HEADER = ["timestamp", "record_id", "detail"]
LEGACY_WEBSITE_HEADER = ["timestamp", "company_name", "proposed_url", "verdict", "detail"]
LEGACY_FIELD_HOLD_HEADER = [
    "timestamp",
    "record_id",
    "company_name",
    "fields_left_uncleared",
    "fields_held_by_website_check",
    "identity_field_renamed",
]

# --- legacy/aekovera/config.py sheet names (verbatim) ----------------------
LEGACY_SHEET_ACCEPTED = "Accepted"
LEGACY_SHEET_REJECTED = "Rejected"
LEGACY_SHEET_UNDECIDED = "Manual review & held"
LEGACY_SHEET_SNAPSHOT = "Accepted companies"


def advancing_clock(base: datetime | None = None):
    """Clock that returns later timestamps on every call."""
    state = {"n": 0}
    base = base or datetime(2026, 9, 21, 12, 0, 0, tzinfo=UTC)

    def _clock() -> datetime:
        state["n"] += 1
        return base.replace(second=min(base.second + state["n"], 59))

    return _clock


def record(record_id: str = "M-1001", name: str = "Alekovera Spice Co") -> dict:
    return {
        "record_id": record_id,
        "fields": {"company_name": name, "city": "Portland", "website_url": "https://x.test"},
    }


def result(decision: str = "accept", **overrides) -> dict:
    base = {
        "decision": decision,
        "company_name": "Alekovera Spice Co",
        "confidence": "high",
        "reason": "in scope",
        "scope_match": "yes",
        "food_beverage_connection": "ingredients for hot sauce",
        "qualifying_supplier_types": ["ingredients"],
        "supply_country": "USA",
        "is_us_based": True,
        "supply_origin_note": "",
        "manual_review_reason": "",
        "changes": [{"field": "city", "new_value": "Portland, OR", "old_value": ""}],
    }
    base.update(overrides)
    return base


def populated_store(tmp_path: Path) -> ReviewStore:
    """A store with accepted, rejected, and held history plus every CSV source."""
    store = ReviewStore(tmp_path / "review.db", clock=advancing_clock())
    store.record_history(
        record("M-1001"),
        result("accept"),
        outcome="accepted",
        mode="auto",
        backend="manual",
        finalized=True,
        applied=["city"],
    )
    store.record_history(
        record("M-1002", "Rejected Foods Inc"),
        result("reject"),
        outcome="rejected",
        mode="auto",
        backend="manual",
        finalized=True,
    )
    store.record_history(
        record("M-1003", "Held Foods"),
        result("accept"),
        outcome="held for field review",
        mode="auto",
        backend="manual",
        finalized=False,
    )
    store.record_manual_review(
        "run-1",
        "M-1004",
        "Scope Unknown LLC",
        reason="scope uncertain",
        food_beverage_connection="packaging only",
        qualifying_supplier_types=["packaging"],
    )
    store.record_field_hold(
        "run-1",
        "M-1005",
        "Held Foods",
        needs_clear=[("products", "???")],
        needs_review=[("website_url", "https://x.test", "no trace of company")],
        identity_renamed=[("company_name", "Held Foods Inc", "Held Foods")],
    )
    store.record_website_check(
        "run-1", "M-1001", "Alekovera Spice Co", "https://x.test", False, "MISMATCH: site sells shoes"
    )
    store.record_website_check("run-1", "M-1002", "Rejected Foods Inc", "https://y.test", None, "timeout")
    store.record_discovery_failure("run-1", "", "no record id found after 3 retries")

    rec = record("M-1001")
    corrected = {
        "fields": {"company_name": "Alekovera Spice Co", "city": "Portland, OR", "website_url": "https://x.test"},
        "record_id": "M-1001",
    }
    store.record_accepted_snapshot(
        rec, corrected, result("accept"), applied=["city"], confirmed_by="auto:verified"
    )
    # A non-US accepted company for accepted_non_us.csv.
    foreign = record("M-1006", "Foreign Spice Traders")
    foreign_result = result(
        "accept",
        company_name="Foreign Spice Traders",
        is_us_based=False,
        supply_country="India",
        supply_origin_note="spices sourced from Kerala farms",
        qualifying_supplier_types=["ingredients", "spices"],
    )
    store.record_accepted_snapshot(
        foreign,
        {"fields": {"company_name": "Foreign Spice Traders"}, "record_id": "M-1006"},
        foreign_result,
        confirmed_by="auto:verified",
    )
    return store


# --------------------------------------------------------------------------- #
# supplier_history.xlsx shape
# --------------------------------------------------------------------------- #
def test_history_meta_block_matches_legacy_verbatim(tmp_path: Path) -> None:
    sheets = history_export(populated_store(tmp_path).history_rows())
    header = sheets[LEGACY_SHEET_ACCEPTED][0]
    assert header[: len(LEGACY_HISTORY_META_COLUMNS)] == LEGACY_HISTORY_META_COLUMNS


def test_history_has_exactly_the_three_legacy_sheets(tmp_path: Path) -> None:
    sheets = history_export(populated_store(tmp_path).history_rows())
    assert set(sheets) == {
        LEGACY_SHEET_ACCEPTED,
        LEGACY_SHEET_REJECTED,
        LEGACY_SHEET_UNDECIDED,
    }


def test_history_detail_columns_follow_legacy_rule(tmp_path: Path) -> None:
    """Preferred order for known fields, then extras sorted - legacy
    ``_ensure_columns`` sort_key, computed per sheet."""
    sheets = history_export(populated_store(tmp_path).history_rows())
    accepted_header = sheets[LEGACY_SHEET_ACCEPTED][0]
    details = accepted_header[len(LEGACY_HISTORY_META_COLUMNS) :]
    # Known detail fields present in the data, in PREFERRED_DETAIL_ORDER...
    assert details[:2] == ["website_url", "city"]
    # ...and the outcome routing column never prints.
    assert "outcome_view" not in details


def test_history_rows_route_to_their_own_sheet(tmp_path: Path) -> None:
    store = populated_store(tmp_path)
    sheets = history_export(store.history_rows())
    # M-1001 (accepted) -> Accepted; M-1006 is snapshot-only (no history row
    # by design - legacy snapshots and history are separate surfaces).
    assert len(sheets[LEGACY_SHEET_ACCEPTED][1]) == 1
    assert len(sheets[LEGACY_SHEET_REJECTED][1]) == 1
    assert len(sheets[LEGACY_SHEET_UNDECIDED][1]) == 1


def test_history_workbook_file_has_legacy_sheet_names_and_headers(tmp_path: Path) -> None:
    store = populated_store(tmp_path)
    out = write_history_workbook(store, tmp_path / "exports" / "supplier_history.xlsx")
    workbook = load_workbook(out)
    assert workbook.sheetnames == [
        LEGACY_SHEET_ACCEPTED,
        LEGACY_SHEET_REJECTED,
        LEGACY_SHEET_UNDECIDED,
    ]
    header = [c.value for c in workbook[LEGACY_SHEET_ACCEPTED][1]]
    assert header[: len(LEGACY_HISTORY_META_COLUMNS)] == LEGACY_HISTORY_META_COLUMNS


# --------------------------------------------------------------------------- #
# accepted_companies.xlsx shape
# --------------------------------------------------------------------------- #
def test_accepted_layout_is_lead_details_tail(tmp_path: Path) -> None:
    header, _ = accepted_export(populated_store(tmp_path).accepted_rows())
    assert header[:3] == LEGACY_ACCEPTED_LEAD_COLUMNS
    assert header[-len(LEGACY_ACCEPTED_TAIL_COLUMNS) :] == LEGACY_ACCEPTED_TAIL_COLUMNS
    # Details follow the legacy PREFERRED_DETAIL_ORDER (accepted_snapshots.py
    # sort_key), not alphabetical: website_url precedes city.
    assert header[3:-len(LEGACY_ACCEPTED_TAIL_COLUMNS)] == ["website_url", "city"]


def test_accepted_workbook_file_shape(tmp_path: Path) -> None:
    store = populated_store(tmp_path)
    out = write_accepted_workbook(store, tmp_path / "exports" / "accepted_companies.xlsx")
    workbook = load_workbook(out)
    assert workbook.sheetnames == [LEGACY_SHEET_SNAPSHOT]
    header = [c.value for c in workbook[LEGACY_SHEET_SNAPSHOT][1]]
    assert header[:3] == LEGACY_ACCEPTED_LEAD_COLUMNS
    assert header[-len(LEGACY_ACCEPTED_TAIL_COLUMNS) :] == LEGACY_ACCEPTED_TAIL_COLUMNS
    # Row values land under the right columns: the corrected city wins.
    rows = list(workbook[LEGACY_SHEET_SNAPSHOT].values)
    by_name = {row[2]: row for row in rows[1:]}
    assert by_name["Alekovera Spice Co"][header.index("city")] == "Portland, OR"


# --------------------------------------------------------------------------- #
# CSV log shapes (all five legacy logs)
# --------------------------------------------------------------------------- #
def test_website_verify_csv_shape(tmp_path: Path) -> None:
    out = export_website_checks(populated_store(tmp_path), tmp_path / "website_verify_log.csv")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0] == ",".join(LEGACY_WEBSITE_HEADER)
    assert lines[1].endswith(",MISMATCH,MISMATCH: site sells shoes")  # False -> MISMATCH
    assert ",inconclusive,timeout" in lines[2]  # None -> inconclusive


def test_manual_review_csv_shape(tmp_path: Path) -> None:
    out = export_manual_reviews(populated_store(tmp_path), tmp_path / "manual_review_queue.csv")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0] == ",".join(LEGACY_MANUAL_REVIEW_HEADER)
    assert "scope uncertain" in lines[1]
    assert "packaging only" in lines[1]
    assert lines[1].endswith("packaging")  # types joined with ", " (single item -> bare)


def test_field_hold_csv_shape(tmp_path: Path) -> None:
    out = export_field_holds(populated_store(tmp_path), tmp_path / "held_for_field_review.csv")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0] == ",".join(LEGACY_FIELD_HOLD_HEADER)
    # Legacy text blocks: f=v / f=v (detail) / f: 'old' -> 'new' (csv.writer
    # quotes only when a comma is present - single-item blocks stay bare).
    assert "products='???'" in lines[1]
    assert "website_url='https://x.test' (no trace of company)" in lines[1]
    assert "company_name: 'Held Foods Inc' -> 'Held Foods'" in lines[1]


def test_discovery_failure_csv_shape_and_unknown_id(tmp_path: Path) -> None:
    out = export_discovery_failures(
        populated_store(tmp_path), tmp_path / "field_discovery_failures.csv"
    )
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0] == ",".join(LEGACY_DISCOVERY_HEADER)
    # Blank record id renders the legacy placeholder.
    assert "(unknown - discovery failed before the id could be read)" in lines[1]


def test_non_us_origin_csv_shape_and_filter(tmp_path: Path) -> None:
    out = export_non_us_origins(populated_store(tmp_path), tmp_path / "accepted_non_us.csv")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0] == ",".join(LEGACY_NON_US_HEADER)
    assert len(lines) == 2  # only the confirmed non-US company (US/None never logged)
    assert "Foreign Spice Traders" in lines[1]
    assert "India" in lines[1]
    assert lines[1].endswith('"ingredients, spices"')  # csv quotes the comma
