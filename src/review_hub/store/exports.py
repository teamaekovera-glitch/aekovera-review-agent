"""v34-parity exports: supplier_history.xlsx, accepted_companies.xlsx, CSV logs.

The store is the system of record; these generators are how an operator keeps
diffing workbooks/CSVs the way they always have. Parity bar: an operator
diffing a v34 workbook against a store export sees the same columns in the
same order. The v34 shapes are derived from the frozen archive
(``legacy/aekovera/history.py`` / ``accepted_snapshots.py`` for the
workbooks, the five CSV writer blocks in ``main.py`` for the logs).

Deliberate deviations (none change a column layout):
- timestamps are the STORE's stored UTC values, not re-stamped local time;
- CSV exports always write a header (legacy only wrote one as a side effect
  of the first log row, so an empty legacy CSV had no header).
Detail columns follow legacy ``_ensure_columns`` exactly: preferred order for
known fields, then anything else sorted, computed per sheet.
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from review_hub.store.repository import ReviewStore
from review_hub.store.rows import (
    ACCEPTED_LEAD_COLUMNS,
    ACCEPTED_TAIL_COLUMNS,
    HISTORY_ACCEPTED_SHEET,
    HISTORY_META_COLUMNS,
    HISTORY_REJECTED_SHEET,
    HISTORY_UNDECIDED_SHEET,
    PREFERRED_DETAIL_ORDER,
    format_pairs,
)

# Legacy log filenames (config.py) - the export defaults.
NON_US_LOG = "accepted_non_us.csv"
MANUAL_REVIEW_LOG = "manual_review_queue.csv"
FIELD_DISCOVERY_FAILURE_LOG = "field_discovery_failures.csv"
WEBSITE_VERIFY_LOG = "website_verify_log.csv"
FIELD_HOLD_LOG = "held_for_field_review.csv"

# Legacy main.py verdict text for the independent website check.
WEBSITE_VERDICT_TEXT = {True: "match", False: "MISMATCH", None: "inconclusive"}

DISCOVERY_UNKNOWN_ID = "(unknown - discovery failed before the id could be read)"

# History sheets in legacy workbook order.
HISTORY_SHEETS = (HISTORY_ACCEPTED_SHEET, HISTORY_REJECTED_SHEET, HISTORY_UNDECIDED_SHEET)


def _detail_columns(rows: Iterable[Mapping[str, Any]], fixed: Sequence[str]) -> list[str]:
    """Company-detail columns for a sheet: legacy preferred order first, then
    any other field appended (sorted for determinism). Mirrors the v34 rule
    that detail columns are discovered from the data, never whitelisted.
    Store bookkeeping never prints: ``_``-prefixed keys and the history
    ``outcome_view`` routing column."""
    present: set[str] = set()
    for row in rows:
        present.update(
            k for k in row if not k.startswith("_") and k != "outcome_view" and k not in fixed
        )
    preferred = [c for c in PREFERRED_DETAIL_ORDER if c in present]
    extras = sorted(present - set(preferred))
    return preferred + extras


def history_export(
    history_rows: Iterable[Mapping[str, Any]],
) -> dict[str, tuple[list[str], list[list[Any]]]]:
    """supplier_history.xlsx as {sheet: (header, rows)}.

    Legacy header rule (history.py ``_ensure_columns``): the meta block
    first, then company-detail columns discovered from the data - preferred
    order for known fields, then anything else sorted - computed PER SHEET
    from that sheet's own rows, so a detail field only ever widens the
    sheets whose suppliers actually carried it. ``outcome_view`` - which
    sheet a row lives in - is a stored column, not a printed one.
    """
    fixed = HISTORY_META_COLUMNS
    rows = list(history_rows)
    sheets: dict[str, tuple[list[str], list[list[Any]]]] = {}
    for name in HISTORY_SHEETS:
        sheet_rows = [row for row in rows if (row.get("outcome_view") or HISTORY_UNDECIDED_SHEET) == name]
        header = [*fixed, *_detail_columns(sheet_rows, fixed)]
        sheets[name] = (header, [[row.get(c, "") for c in header] for row in sheet_rows])
    return sheets


def accepted_export(
    accepted_rows: Iterable[Mapping[str, Any]],
) -> tuple[list[str], list[list[Any]]]:
    """accepted_companies.xlsx ("Accepted companies") as (header, rows).

    Legacy layout: identity (accepted_at, record_id, company_name), then the
    company's details, then the edit-summary tail.
    """
    rows = list(accepted_rows)
    fixed = [*ACCEPTED_LEAD_COLUMNS, *ACCEPTED_TAIL_COLUMNS]
    details = _detail_columns(rows, fixed)
    header = [*ACCEPTED_LEAD_COLUMNS, *details, *ACCEPTED_TAIL_COLUMNS]
    return header, [[row.get(c, "") for c in header] for row in rows]


# Legacy history.py _style: wide-text columns get 42, everything else is
# fitted from the header itself.
_HISTORY_WIDE_COLUMNS = {
    "reason",
    "food_beverage_connection",
    "manual_review_reason",
    "products",
    "description",
    "specialty",
    "supply_origin_note",
}
_HEADER_FILL = "FFEFEFEF"


def _style_sheet(worksheet: Any) -> None:
    """Freeze the header row, bold it, add a filter and sane column widths
    (ported from legacy ``history._style``) so the export reads like the
    workbook the operator already knows."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if worksheet.max_row < 1:
        return
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"
    try:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(worksheet.max_column)}1"
    except Exception:  # pragma: no cover - openpyxl filter edge cases
        pass
    wide = _HISTORY_WIDE_COLUMNS | {"edits_made"}
    for col in range(1, worksheet.max_column + 1):
        header = str(worksheet.cell(row=1, column=col).value or "")
        width = 42 if header in wide else max(12, min(len(header) + 6, 28))
        worksheet.column_dimensions[get_column_letter(col)].width = width


def write_history_workbook(store: ReviewStore, path: str | Path) -> Path:
    """Write supplier_history.xlsx (3 sheets) from the store."""
    sheets = history_export(store.history_rows())
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in HISTORY_SHEETS:
        header, rows = sheets[sheet_name]
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(header)
        for row in rows:
            worksheet.append(row)
        _style_sheet(worksheet)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return out


def write_accepted_workbook(store: ReviewStore, path: str | Path) -> Path:
    """Write accepted_companies.xlsx ("Accepted companies") from the store."""
    header, rows = accepted_export(store.accepted_rows())
    workbook = Workbook()
    workbook.remove(workbook.active)
    worksheet = workbook.create_sheet("Accepted companies")
    worksheet.append(header)
    for row in rows:
        worksheet.append(row)
    _style_sheet(worksheet)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return out


def write_csv(path: str | Path, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> Path:
    """Write one CSV log with a v34 header shape (header always written)."""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            writer.writerow(row)
    return out


def website_check_csv_rows(events: Iterable[Mapping[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """website_verify_log.csv shape: timestamp, company_name, proposed_url,
    verdict (match / MISMATCH / inconclusive), detail."""
    header = ["timestamp", "company_name", "proposed_url", "verdict", "detail"]
    rows = []
    for event in events:
        payload = event.get("payload") or {}
        verdict = payload.get("verdict")
        rows.append(
            [
                event.get("created_at", ""),
                payload.get("company_name", ""),
                payload.get("proposed_url", ""),
                WEBSITE_VERDICT_TEXT.get(verdict, WEBSITE_VERDICT_TEXT[None]),
                event.get("detail", ""),
            ]
        )
    return header, rows


def export_website_checks(store: ReviewStore, path: str | Path) -> Path:
    events = [e for e in store.audit_events() if e.get("kind") == "website_check"]
    header, rows = website_check_csv_rows(events)
    return write_csv(path, header, rows)


def manual_review_csv_rows(holds: Iterable[Mapping[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """manual_review_queue.csv shape."""
    header = [
        "timestamp",
        "record_id",
        "company_name",
        "manual_review_reason",
        "food_beverage_connection",
        "qualifying_supplier_types",
    ]
    rows = []
    for hold in holds:
        detail = hold.get("detail") or {}
        rows.append(
            [
                hold.get("created_at", ""),
                hold.get("record_id", ""),
                hold.get("company_name", ""),
                detail.get("manual_review_reason", ""),
                detail.get("food_beverage_connection", ""),
                ", ".join(detail.get("qualifying_supplier_types") or []),
            ]
        )
    return header, rows


def export_manual_reviews(store: ReviewStore, path: str | Path) -> Path:
    holds = [h for h in store.holds() if h.get("kind") == "manual_review"]
    header, rows = manual_review_csv_rows(holds)
    return write_csv(path, header, rows)


def field_hold_csv_rows(holds: Iterable[Mapping[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """held_for_field_review.csv shape: the legacy text blocks
    (``f=v`` / ``f=v (detail)`` / ``f: 'old' -> 'new'``), joined with ``; ``."""
    header = [
        "timestamp",
        "record_id",
        "company_name",
        "fields_left_uncleared",
        "fields_held_by_website_check",
        "identity_field_renamed",
    ]
    rows = []
    for hold in holds:
        detail = hold.get("detail") or {}
        rows.append(
            [
                hold.get("created_at", ""),
                hold.get("record_id", ""),
                hold.get("company_name", ""),
                format_pairs(detail.get("needs_clear") or (), lambda x: f"{x[0]}={x[1]!r}"),
                format_pairs(
                    detail.get("needs_review") or (), lambda x: f"{x[0]}={x[1]!r} ({x[2]})"
                ),
                format_pairs(
                    detail.get("identity_renamed") or (),
                    lambda x: f"{x[0]}: {x[1]!r} -> {x[2]!r}",
                ),
            ]
        )
    return header, rows


def export_field_holds(store: ReviewStore, path: str | Path) -> Path:
    holds = [h for h in store.holds() if h.get("kind") == "field_hold"]
    header, rows = field_hold_csv_rows(holds)
    return write_csv(path, header, rows)


def discovery_failure_csv_rows(
    events: Iterable[Mapping[str, Any]],
) -> tuple[list[str], list[list[Any]]]:
    """field_discovery_failures.csv shape."""
    header = ["timestamp", "record_id", "detail"]
    rows = []
    for event in events:
        record_id = event.get("record_id", "")
        rows.append(
            [
                event.get("created_at", ""),
                record_id if record_id else DISCOVERY_UNKNOWN_ID,
                event.get("detail", ""),
            ]
        )
    return header, rows


def export_discovery_failures(store: ReviewStore, path: str | Path) -> Path:
    events = [e for e in store.audit_events() if e.get("kind") == "discovery_failure"]
    header, rows = discovery_failure_csv_rows(events)
    return write_csv(path, header, rows)


def non_us_csv_rows(accepted_rows: Iterable[Mapping[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """accepted_non_us.csv shape: one row per accepted company the engine
    confirmed is NOT US-based (True/None never logged - legacy
    ``record_supply_origin``)."""
    header = [
        "timestamp",
        "company_name",
        "supply_country",
        "supply_origin_note",
        "supplier_types",
    ]
    rows = []
    for row in accepted_rows:
        decision = row.get("_decision") or {}
        if decision.get("is_us_based") is not False:
            continue
        rows.append(
            [
                row.get("accepted_at", ""),
                row.get("company_name", ""),
                decision.get("supply_country", ""),
                decision.get("supply_origin_note", ""),
                ", ".join(decision.get("qualifying_supplier_types") or []),
            ]
        )
    return header, rows


def export_non_us_origins(store: ReviewStore, path: str | Path) -> Path:
    header, rows = non_us_csv_rows(store.accepted_rows())
    return write_csv(path, header, rows)
